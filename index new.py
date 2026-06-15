import pandas as pd
import ttkbootstrap as ttk
from ttkbootstrap.constants import *
from tkinter import filedialog, messagebox, BooleanVar
import tkinter as tk
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os
import json
import webbrowser
from datetime import datetime
import math
import random
import statistics


# ══════════════════════════════════════════════════════════════════════════════
#                           STYLES EXCEL
# ══════════════════════════════════════════════════════════════════════════════

class StylesExcel:
    BORDURE = Border(
        left=Side(style='thin', color="D5D8DC"),
        right=Side(style='thin', color="D5D8DC"),
        top=Side(style='thin', color="D5D8DC"),
        bottom=Side(style='thin', color="D5D8DC")
    )

    HEADER_FILL = PatternFill(start_color="27AE60", end_color="27AE60", fill_type="solid")
    HEADER_FONT = Font(bold=True, color="FFFFFF", size=11, name="Calibri")
    HEADER_ALIGN = Alignment(horizontal="center", vertical="center")
    HEADER_VERTICAL = Alignment(textRotation=90, horizontal="center", vertical="center")

    TITRE_FICHIER_COLORS = [
        "FFE082", "64B5F6", "81C784", "BA68C8", "FF8A65",
        "4DD0E1", "FFD54F", "9575CD", "AED581", "FFB74D",
        "F48FB1", "90CAF9", "A5D6A7", "CE93D8", "FFAB91",
        "80DEEA", "FFF59D", "B39DDB", "C5E1A5", "FFCC80",
        "EF9A9A", "81D4FA", "C8E6C9", "E1BEE7", "FFCCBC",
        "B2DFDB", "FFF9C4", "D1C4E9", "DCEDC8", "FFE0B2",
        "F8BBD0", "BBDEFB", "DCEDC8", "F3E5F5", "FFE0B2",
        "FFCDD2", "B3E5FC", "E8F5E9", "F3E5F5", "FBE9E7"
    ]

    HEADER_COLORS = [
        "FFFED9", "FFEAE9", "E3FCE7", "DCD7FE", "FCD9F2",
        "FFE1C2", "F2F1F0", "E6E2D3", "FEFFD6", "FAD3A8",
        "E8EAF6", "FCE4EC", "F1F8E9", "FFF3E0", "E0F2F1",
        "FFF8E1", "F3E5F5", "E8F5E9", "FBE9E7", "ECEFF1",
        "E1F5FE", "F9FBE7", "FCE4EC", "EFEBE9", "FFF3E0",
        "EDE7F6", "F1F8E9", "E0F7FA", "FFFDE7", "FFEBEE",
        "E8EAF6", "E0F2F1", "FFF9C4", "F8BBD0", "DCEDC8",
        "B2EBF2", "FFF59D", "E1BEE7", "C5E1A5", "FFCCBC"
    ]

    RANDOM_TITRE_FILL = None
    RANDOM_HEADER_FILL = None

    TITRE_ASSEMBLEUR_FONT = Font(color="000000", bold=True, size=11, name="Calibri")
    HEADER_ASSEMBLEUR_FONT = Font(bold=True, color="000000", size=11, name="Calibri")
    TITRE_ALIGN = Alignment(horizontal="center", vertical="center")

    VERT = PatternFill(start_color="D5F5E3", end_color="D5F5E3", fill_type="solid")
    ROUGE = PatternFill(start_color="FADBD8", end_color="FADBD8", fill_type="solid")
    JAUNE = PatternFill(start_color="FCF3CF", end_color="FCF3CF", fill_type="solid")
    BLEU_A = PatternFill(start_color="D6EAF8", end_color="D6EAF8", fill_type="solid")
    ORANGE_ST = PatternFill(start_color="FDEBD0", end_color="FDEBD0", fill_type="solid")

    LIGNE_1 = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    LIGNE_2 = PatternFill(start_color="F8F9FA", end_color="F8F9FA", fill_type="solid")

    DATA_FONT = Font(size=10, name="Calibri", color="2C3E50")
    GRAS_FONT = Font(bold=True, size=11, name="Calibri", color="1F4E79")
    CENTER_ALIGN = Alignment(horizontal="center", vertical="center")

    @classmethod
    def init_random_assembleur_colors(cls):
        titre_color = random.choice(cls.TITRE_FICHIER_COLORS)
        header_color = random.choice(cls.HEADER_COLORS)
        cls.RANDOM_TITRE_FILL = PatternFill(
            start_color=titre_color, end_color=titre_color, fill_type="solid"
        )
        cls.RANDOM_HEADER_FILL = PatternFill(
            start_color=header_color, end_color=header_color, fill_type="solid"
        )

    @staticmethod
    def appliquer_style_moderne(ws, titre_rows=None):
        if titre_rows is None:
            titre_rows = []
        ws.sheet_view.showGridLines = False
        for row_idx in range(2, ws.max_row + 1):
            if row_idx in titre_rows:
                continue
            premiere_cellule = ws.cell(row=row_idx, column=1)
            is_special = False
            if premiere_cellule.value:
                val = str(premiere_cellule.value).strip().upper()
                if val in ["A", "ST", "INTERNE", "SOUS TRAITANT"] or "TOTAL" in val:
                    is_special = True
            for col_idx in range(1, ws.max_column + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                if not is_special:
                    try:
                        current_color = cell.fill.start_color.rgb
                    except Exception:
                        current_color = None
                    if current_color in [None, "00000000", "FFFFFF", "00FFFFFF"]:
                        cell.fill = (
                            StylesExcel.LIGNE_1 if row_idx % 2 == 0
                            else StylesExcel.LIGNE_2
                        )
                cell.border = StylesExcel.BORDURE
                if not cell.alignment.wrap_text and cell.alignment.text_rotation != 90:
                    cell.alignment = StylesExcel.CENTER_ALIGN


# ══════════════════════════════════════════════════════════════════════════════
#                        MISE EN FORME
# ══════════════════════════════════════════════════════════════════════════════

class MiseEnForme:
    @staticmethod
    def appliquer(fichier):
        wb = load_workbook(fichier)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            header_rows = []
            titre_rows = []
            VALEURS_RESERVEES = {'INTERNE', 'SOUS TRAITANT', 'A', 'ST', 'TOTAL'}
            for row_idx in range(1, ws.max_row + 1):
                cell = ws.cell(row=row_idx, column=1)
                if cell.value and isinstance(cell.value, str):
                    val_upper = cell.value.strip().upper()
                    if val_upper in VALEURS_RESERVEES:
                        continue
                    row_values = [
                        ws.cell(row=row_idx, column=c).value
                        for c in range(2, ws.max_column + 1)
                    ]
                    if not any(row_values):
                        titre_rows.append(row_idx)
                        header_rows.append(row_idx + 1)
                        if ws.max_column >= 1:
                            try:
                                ws.merge_cells(
                                    start_row=row_idx, start_column=1,
                                    end_row=row_idx, end_column=ws.max_column
                                )
                            except Exception:
                                pass
                        titre_fill = StylesExcel.RANDOM_TITRE_FILL
                        cell.fill = titre_fill
                        cell.font = StylesExcel.TITRE_ASSEMBLEUR_FONT
                        cell.alignment = StylesExcel.TITRE_ALIGN
                        cell.border = StylesExcel.BORDURE

            if sheet_name == "RAMASSAGE PAR PRESTATAIRE":
                MiseEnForme._formater_par_p(ws, header_rows)
            elif sheet_name == "A et ST":
                MiseEnForme._formater_a_et_st(ws, header_rows)
            elif sheet_name in ["RAMASSAGE PAR JOUR", "RAMASSAGE PAR SEMAINE", "RAMASSAGE PAR MOIS"]:
                MiseEnForme._formater_periode(ws, header_rows)
            elif sheet_name in ["RAMASSAGE PAR MODELES", "DISPATCH"]:
                MiseEnForme._formater_reference_dispatch(ws, header_rows)
            elif sheet_name == "STATS_TEMPS":
                MiseEnForme._formater_stats_temps(ws, header_rows)
            elif sheet_name == "TIMING PAR MODELE":
                MiseEnForme._formater_timing(ws, header_rows)
            else:
                MiseEnForme._formater_autres(ws, header_rows)

            MiseEnForme._colorer_total(ws, titre_rows)
            MiseEnForme._optimiser_hauteur_lignes(ws, titre_rows)
            StylesExcel.appliquer_style_moderne(ws, titre_rows)
        wb.save(fichier)

    @staticmethod
    def _optimiser_hauteur_lignes(ws, titre_rows):
        for row in range(1, ws.max_row + 1):
            if row in titre_rows:
                ws.row_dimensions[row].height = 18
            else:
                ws.row_dimensions[row].height = None

    @staticmethod
    def _formater_timing(ws, header_rows):
        header_fill = StylesExcel.RANDOM_HEADER_FILL
        header_font = StylesExcel.HEADER_ASSEMBLEUR_FONT
        for hr in header_rows:
            if hr > ws.max_row:
                continue
            for col in range(1, ws.max_column + 1):
                cell = ws.cell(row=hr, column=col)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = StylesExcel.HEADER_ALIGN
                cell.border = StylesExcel.BORDURE
        ws.column_dimensions['A'].width = 25
        ws.column_dimensions['B'].width = 12
        for col in range(3, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(col)].width = 16

    @staticmethod
    def _formater_stats_temps(ws, header_rows):
        header_fill = StylesExcel.RANDOM_HEADER_FILL
        header_font = StylesExcel.HEADER_ASSEMBLEUR_FONT
        for hr in header_rows:
            if hr > ws.max_row:
                continue
            for col in range(1, ws.max_column + 1):
                cell = ws.cell(row=hr, column=col)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = StylesExcel.HEADER_ALIGN
                cell.border = StylesExcel.BORDURE
        ws.column_dimensions['A'].width = 25
        ws.column_dimensions['B'].width = 12
        for col in range(3, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(col)].width = 18

    @staticmethod
    def _formater_par_p(ws, header_rows):
        header_fill = StylesExcel.RANDOM_HEADER_FILL
        header_font = StylesExcel.HEADER_ASSEMBLEUR_FONT
        for hr in header_rows:
            if hr > ws.max_row:
                continue
            headers = [
                str(ws.cell(row=hr, column=c).value or "").upper()
                for c in range(1, ws.max_column + 1)
            ]
            col_details = next(
                (i + 1 for i, h in enumerate(headers) if "DETAILS" in h), None
            )
            if col_details:
                for row in range(hr + 1, ws.max_row + 1):
                    cell = ws.cell(row=row, column=col_details)
                    if cell.value and isinstance(cell.value, str):
                        cell.value = str(cell.value).replace(', ', '\n')
                        cell.alignment = Alignment(
                            wrap_text=True, vertical="top", horizontal="left"
                        )
            for col in range(1, ws.max_column + 1):
                cell = ws.cell(row=hr, column=col)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = StylesExcel.HEADER_VERTICAL
                cell.border = StylesExcel.BORDURE
        ws.column_dimensions['A'].width = 10
        ws.column_dimensions['B'].width = 36
        ws.column_dimensions['C'].width = 13
        for col in range(4, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(col)].width = 5.5

    @staticmethod
    def _formater_a_et_st(ws, header_rows):
        header_fill = StylesExcel.RANDOM_HEADER_FILL
        header_font = StylesExcel.HEADER_ASSEMBLEUR_FONT
        for row in ws.iter_rows():
            for cell in row:
                if cell.value in ['INTERNE', 'SOUS TRAITANT']:
                    cell.fill = StylesExcel.JAUNE
                    cell.font = StylesExcel.GRAS_FONT
        for hr in header_rows:
            if hr > ws.max_row:
                continue
            for col in range(1, ws.max_column + 1):
                cell = ws.cell(row=hr, column=col)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = StylesExcel.HEADER_ALIGN
                cell.border = StylesExcel.BORDURE
        for col in ws.columns:
            try:
                letter = get_column_letter(col[0].column)
                max_len = max(
                    (len(str(c.value)) for c in col if c.value is not None),
                    default=8
                )
                ws.column_dimensions[letter].width = min(max_len + 3, 50)
            except Exception:
                continue

    @staticmethod
    def _formater_periode(ws, header_rows):
        header_fill = StylesExcel.RANDOM_HEADER_FILL
        header_font = StylesExcel.HEADER_ASSEMBLEUR_FONT
        for hr in header_rows:
            if hr > ws.max_row:
                continue
            for col in range(1, ws.max_column + 1):
                cell = ws.cell(row=hr, column=col)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = StylesExcel.HEADER_VERTICAL
                cell.border = StylesExcel.BORDURE
        for col in range(1, ws.max_column + 1):
            if col == 1:
                max_len = max(
                    (len(str(c.value))
                     for c in ws[get_column_letter(col)] if c.value),
                    default=0
                )
                ws.column_dimensions[get_column_letter(col)].width = min(max_len + 2, 50)
            else:
                ws.column_dimensions[get_column_letter(col)].width = 5.5
        for row in range(1, ws.max_row + 1):
            val = str(ws.cell(row=row, column=1).value or "").strip().upper()
            if val == "A":
                for col in range(1, ws.max_column + 1):
                    ws.cell(row=row, column=col).fill = StylesExcel.BLEU_A
            elif val == "ST":
                for col in range(1, ws.max_column + 1):
                    ws.cell(row=row, column=col).fill = StylesExcel.ORANGE_ST

    @staticmethod
    def _formater_reference_dispatch(ws, header_rows):
        header_fill = StylesExcel.RANDOM_HEADER_FILL
        header_font = StylesExcel.HEADER_ASSEMBLEUR_FONT
        headers = [str(c.value or "") for c in ws[1]]
        date_cols = [i + 1 for i, h in enumerate(headers) if "/" in h]
        col_reste_livrer = None
        for i, h in enumerate(headers):
            h_upper = h.upper()
            if "RESTE" in h_upper and "LIVRER" in h_upper:
                col_reste_livrer = i + 1
                break
        for hr in header_rows:
            if hr > ws.max_row:
                continue
            for col in range(1, ws.max_column + 1):
                cell = ws.cell(row=hr, column=col)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = StylesExcel.HEADER_VERTICAL
                cell.border = StylesExcel.BORDURE
        for col in range(1, ws.max_column + 1):
            if col == 1:
                max_len = max(
                    (len(str(c.value))
                     for c in ws[get_column_letter(col)] if c.value),
                    default=0
                )
                ws.column_dimensions[get_column_letter(col)].width = min(max_len + 2, 50)
            else:
                ws.column_dimensions[get_column_letter(col)].width = 5.5
        for row in range(2, ws.max_row + 1):
            for col in date_cols:
                cell = ws.cell(row=row, column=col)
                try:
                    if cell.value and float(cell.value) != 0:
                        cell.fill = StylesExcel.VERT
                except Exception:
                    pass
            if col_reste_livrer:
                cell = ws.cell(row=row, column=col_reste_livrer)
                try:
                    if cell.value and float(cell.value) > 0:
                        cell.fill = StylesExcel.ROUGE
                except Exception:
                    pass

    @staticmethod
    def _formater_autres(ws, header_rows):
        header_fill = StylesExcel.RANDOM_HEADER_FILL
        header_font = StylesExcel.HEADER_ASSEMBLEUR_FONT
        for hr in header_rows:
            if hr > ws.max_row:
                continue
            for col in range(1, ws.max_column + 1):
                cell = ws.cell(row=hr, column=col)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = StylesExcel.HEADER_ALIGN
                cell.border = StylesExcel.BORDURE
        for col in ws.columns:
            try:
                letter = get_column_letter(col[0].column)
                max_len = max(
                    (len(str(c.value)) for c in col if c.value is not None),
                    default=8
                )
                ws.column_dimensions[letter].width = min(max_len + 3, 50)
            except Exception:
                continue

    @staticmethod
    def _colorer_total(ws, titre_rows):
        for row in range(1, ws.max_row + 1):
            if row in titre_rows:
                continue
            val = str(ws.cell(row=row, column=1).value or "").strip().upper()
            if "TOTAL" in val:
                couleur = "B8DAFF"
                fill = PatternFill(
                    start_color=couleur, end_color=couleur, fill_type="solid"
                )
                for col in range(1, ws.max_column + 1):
                    cell = ws.cell(row=row, column=col)
                    cell.fill = fill
                    cell.font = StylesExcel.GRAS_FONT
                    cell.alignment = StylesExcel.CENTER_ALIGN


# ══════════════════════════════════════════════════════════════════════════════
#                        LECTEUR RECAP GLOBAL (pour HTML)
# ══════════════════════════════════════════════════════════════════════════════

class LecteurRecapGlobal:

    @staticmethod
    def lire_fichier(fichier_excel):
        wb = load_workbook(fichier_excel, data_only=True)
        if "RECAP GLOBAL" not in wb.sheetnames:
            raise ValueError("Feuille 'RECAP GLOBAL' introuvable.")
        ws = wb["RECAP GLOBAL"]
        titre_rows = []
        for row_idx in range(1, ws.max_row + 1):
            cell = ws.cell(row=row_idx, column=1)
            if cell.value and isinstance(cell.value, str):
                row_values = [ws.cell(row=row_idx, column=c).value
                              for c in range(2, ws.max_column + 1)]
                if not any(row_values):
                    titre_rows.append(row_idx)
        blocs = []
        if not titre_rows:
            blocs.append({
                'titre': 'RECAP GLOBAL',
                'df': LecteurRecapGlobal._extraire_bloc(ws, 1, ws.max_row)
            })
        else:
            for i, row_titre in enumerate(titre_rows):
                titre_val = str(ws.cell(row=row_titre, column=1).value or "").strip()
                header_row = row_titre + 1
                end_row = (titre_rows[i + 1] - 1
                           if i + 1 < len(titre_rows) else ws.max_row)
                df = LecteurRecapGlobal._extraire_bloc(ws, header_row, end_row)
                if df is not None and len(df) > 0:
                    blocs.append({'titre': titre_val, 'df': df})
        return blocs

    @staticmethod
    def _extraire_bloc(ws, header_row, end_row):
        if header_row > ws.max_row or header_row > end_row:
            return None
        headers = []
        for c in range(1, ws.max_column + 1):
            val = ws.cell(row=header_row, column=c).value
            headers.append(str(val).strip() if val is not None else f"COL_{c}")
        lignes = []
        for r in range(header_row + 1, end_row + 1):
            ligne = {}
            has_data = False
            for c, h in enumerate(headers, 1):
                val = ws.cell(row=r, column=c).value
                ligne[h] = val
                if val is not None:
                    has_data = True
            if has_data:
                lignes.append(ligne)
        if not lignes:
            return None
        return pd.DataFrame(lignes, columns=headers)

    @staticmethod
    def preparer_donnees_graphe(df):
        df_clean = df.copy()
        df_clean.columns = [str(c).strip().upper() for c in df_clean.columns]
        col_ref = next((c for c in df_clean.columns if 'REFERENCE' in c), None)
        if col_ref is None:
            return None
        mask = df_clean[col_ref].astype(str).str.upper().str.strip() != 'TOTAL'
        df_data = df_clean[mask].copy()
        if len(df_data) == 0:
            return None
        CIBLES = {
            'QTE COMMANDE':     ['QTE COMMANDE', 'QTE', 'COMMANDE'],
            'QTES DISPATCHEES': ['QTES DISPATCHEES', 'DISPATCHEES', 'DISPATCH'],
            'QTES LIVREES':     ['QTES LIVREES', 'LIVREES', 'LIVRAISON'],
            'TOTAL EXP':        ['TOTAL EXP', 'TOTAL_EXP', 'TOTALEXP'],
        }

        def find_col(cands):
            for cand in cands:
                for col in df_clean.columns:
                    if cand.upper() in col.upper():
                        return col
            return None

        res = {k: find_col(v) for k, v in CIBLES.items()}

        def vals(col_name):
            if col_name is None or col_name not in df_data.columns:
                return [0] * len(df_data)
            out = []
            for v in df_data[col_name]:
                try:
                    n = float(v) if v is not None else 0
                    out.append(int(n) if n == int(n) else n)
                except Exception:
                    out.append(0)
            return out

        qte_cmd   = vals(res['QTE COMMANDE'])
        qtes_disp = vals(res['QTES DISPATCHEES'])
        qtes_livr = vals(res['QTES LIVREES'])
        total_exp = vals(res['TOTAL EXP'])
        has_total_exp = res['TOTAL EXP'] is not None

        qte_en_cours = [max(0, qtes_disp[i] - qtes_livr[i]) for i in range(len(qte_cmd))]

        seg_exp = []; seg_livrees = []; seg_dispatch = []; seg_reste = []
        for i in range(len(qte_cmd)):
            cmd = max(qte_cmd[i], 0)
            exp = min(max(total_exp[i], 0), cmd) if has_total_exp else 0
            liv = min(max(qtes_livr[i], 0), cmd)
            dis = min(max(qtes_disp[i], 0), cmd)
            seg_exp.append(exp)
            seg_livrees.append(max(liv - exp, 0))
            seg_dispatch.append(max(dis - liv, 0))
            seg_reste.append(max(cmd - dis, 0))

        return {
            'references':       df_data[col_ref].astype(str).tolist(),
            'qte_commande':     qte_cmd,
            'qtes_dispatchees': qtes_disp,
            'qtes_livrees':     qtes_livr,
            'total_exp':        total_exp,
            'qte_en_cours':     qte_en_cours,
            'seg_exp':          seg_exp,
            'seg_livrees':      seg_livrees,
            'seg_dispatch':     seg_dispatch,
            'seg_reste':        seg_reste,
            'has_total_exp':    has_total_exp,
        }


# ══════════════════════════════════════════════════════════════════════════════
#                         GÉNÉRATEUR HTML
# ══════════════════════════════════════════════════════════════════════════════

class GenerateurHTML:

    @staticmethod
    def generer(blocs, fichier_sortie, titre_page):
        datasets_json = []
        for bloc in blocs:
            donnees = LecteurRecapGlobal.preparer_donnees_graphe(bloc['df'])
            if donnees:
                datasets_json.append({'titre': bloc['titre'], 'donnees': donnees})
        if not datasets_json:
            raise ValueError("Aucune donnée valide pour le graphe.")
        data_js = json.dumps(datasets_json, ensure_ascii=False)
        with open(fichier_sortie, 'w', encoding='utf-8') as f:
            f.write(GenerateurHTML._html(data_js, titre_page))

    @staticmethod
    def _html(data_js, titre_page):
        return f"""<!DOCTYPE html>
<html lang="fr" data-theme="dark">
<head>
<meta charset="UTF-8"/>
<meta name="viewport" content="width=device-width,initial-scale=1"/>
<title>{titre_page}</title>
<link rel="preconnect" href="https://fonts.googleapis.com"/>
<link href="https://fonts.googleapis.com/css2?family=DM+Sans:wght@300;400;500;600;700&family=DM+Mono:wght@400;500&display=swap" rel="stylesheet"/>
<script src="https://cdn.jsdelivr.net/npm/chart.js@4.4.0/dist/chart.umd.min.js"></script>
<style>
:root,[data-theme="dark"]{{
  --bg:#080e1a;--bg2:#0d1526;--bg3:#121e33;--bg4:#182540;
  --border:rgba(255,255,255,0.06);--border2:rgba(255,255,255,0.10);
  --text:#dde4f0;--text2:#7a8fab;--muted:#3d5068;
  --shadow:0 20px 60px rgba(0,0,0,0.7);
  --card-shadow:0 4px 24px rgba(0,0,0,0.5);
  --glow:0 0 40px rgba(20,184,166,0.08);
}}
[data-theme="light"]{{
  --bg:#f0f5fb;--bg2:#ffffff;--bg3:#e8eef7;--bg4:#dde5f0;
  --border:rgba(0,0,0,0.07);--border2:rgba(0,0,0,0.12);
  --text:#1a2540;--text2:#4a6080;--muted:#8fa3be;
  --shadow:0 20px 60px rgba(0,0,0,0.10);
  --card-shadow:0 4px 20px rgba(0,0,0,0.07);
  --glow:0 0 40px rgba(20,184,166,0.05);
}}
*,*::before,*::after{{box-sizing:border-box;margin:0;padding:0}}
html{{scroll-behavior:smooth}}
body{{font-family:'DM Sans',system-ui,sans-serif;background:var(--bg);color:var(--text);min-height:100vh;transition:background .4s,color .4s;}}
::-webkit-scrollbar{{width:4px;height:4px}}
::-webkit-scrollbar-thumb{{background:var(--bg4);border-radius:99px}}
body::before{{
  content:'';position:fixed;top:-200px;left:50%;transform:translateX(-50%);
  width:900px;height:500px;
  background:radial-gradient(ellipse at center,rgba(20,184,166,0.07) 0%,rgba(99,102,241,0.04) 40%,transparent 70%);
  pointer-events:none;z-index:0;transition:opacity .4s;
}}
[data-theme="light"] body::before{{opacity:0}}
header{{
  height:58px;display:flex;align-items:center;gap:16px;padding:0 28px;
  background:rgba(8,14,26,0.85);backdrop-filter:blur(16px);
  border-bottom:1px solid var(--border);position:sticky;top:0;z-index:200;
  transition:background .4s,border-color .4s;
}}
[data-theme="light"] header{{background:rgba(255,255,255,0.85);}}
.logo-mark{{
  width:28px;height:28px;border-radius:8px;
  background:linear-gradient(135deg,#14b8a6,#6366f1);
  display:flex;align-items:center;justify-content:center;
  font-size:.72rem;font-weight:700;color:#fff;flex-shrink:0;
  box-shadow:0 0 16px rgba(20,184,166,0.35);
}}
header h1{{font-size:.92rem;font-weight:600;flex:1;white-space:nowrap;overflow:hidden;text-overflow:ellipsis;color:var(--text);}}
.badge{{font-size:.6rem;font-weight:700;letter-spacing:1.2px;text-transform:uppercase;padding:3px 9px;border-radius:5px;white-space:nowrap;}}
.bd1{{color:#14b8a6;background:rgba(20,184,166,0.12);border:1px solid rgba(20,184,166,0.25);}}
.bd2{{color:#6366f1;background:rgba(99,102,241,0.12);border:1px solid rgba(99,102,241,0.25);}}
.theme-btn{{
  display:flex;align-items:center;gap:8px;padding:6px 13px;border-radius:8px;cursor:pointer;
  border:1px solid var(--border2);background:var(--bg3);color:var(--text2);
  font-size:.75rem;font-weight:500;font-family:inherit;transition:all .2s;white-space:nowrap;user-select:none;
}}
.theme-btn:hover{{border-color:rgba(20,184,166,.4);color:var(--text);background:var(--bg4)}}
.tog-track{{width:32px;height:18px;border-radius:99px;background:var(--bg4);position:relative;border:1px solid var(--border2);transition:background .3s;flex-shrink:0;}}
[data-theme="light"] .tog-track{{background:linear-gradient(90deg,#14b8a6,#6366f1);border-color:transparent;}}
.tog-thumb{{width:12px;height:12px;border-radius:50%;background:#fff;position:absolute;top:2px;left:2px;transition:transform .3s cubic-bezier(.34,1.56,.64,1);box-shadow:0 1px 4px rgba(0,0,0,.4);}}
[data-theme="light"] .tog-thumb{{transform:translateX(14px)}}
.toolbar{{
  background:var(--bg2);border-bottom:1px solid var(--border);
  padding:10px 28px;display:flex;align-items:center;gap:12px;flex-wrap:wrap;
  transition:background .4s,border-color .4s;position:relative;z-index:1;
}}
.tb-lbl{{font-size:.65rem;font-weight:700;color:var(--muted);text-transform:uppercase;letter-spacing:1px;white-space:nowrap;font-family:'DM Mono',monospace;}}
.sw{{position:relative;flex:1;min-width:200px;max-width:440px}}
.sw::after{{content:'›';position:absolute;right:11px;top:50%;transform:translateY(-50%) rotate(90deg);color:#14b8a6;font-size:1.1rem;font-weight:700;pointer-events:none;}}
select#sel{{
  width:100%;padding:7px 30px 7px 12px;background:var(--bg3);border:1px solid var(--border2);
  border-radius:8px;color:var(--text);font-size:.84rem;font-family:'DM Sans',inherit;font-weight:500;
  cursor:pointer;appearance:none;outline:none;transition:border-color .2s,box-shadow .2s,background .4s;
}}
select#sel:hover,select#sel:focus{{border-color:#14b8a6;box-shadow:0 0 0 3px rgba(20,184,166,0.12);}}
.tb-ct{{font-size:.7rem;color:var(--muted);white-space:nowrap;font-family:'DM Mono',monospace;}}
.tb-ct b{{color:#14b8a6;font-weight:700;}}
main{{padding:24px 28px 60px;max-width:1400px;margin:0 auto;position:relative;z-index:1;}}
.gc{{display:none;animation:fadeUp .38s cubic-bezier(.22,.68,0,1.05)}}
.gc.active{{display:block}}
@keyframes fadeUp{{from{{opacity:0;transform:translateY(14px)}}to{{opacity:1;transform:translateY(0)}}}}
.gc-title{{
  font-size:1.15rem;font-weight:700;letter-spacing:-.02em;
  background:linear-gradient(90deg,var(--text) 60%,var(--text2));
  -webkit-background-clip:text;-webkit-text-fill-color:transparent;margin-bottom:18px;
}}
.kpi-row{{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:10px;margin-bottom:20px;}}
.kpi{{
  background:var(--bg2);border:1px solid var(--border);border-radius:12px;
  padding:16px 18px 14px;position:relative;overflow:hidden;cursor:default;
  transition:transform .15s,box-shadow .15s,background .4s,border-color .35s;
}}
.kpi:hover{{transform:translateY(-3px);box-shadow:var(--shadow);border-color:var(--border2);}}
.kpi::before{{content:'';position:absolute;top:0;left:0;right:0;height:2px;border-radius:12px 12px 0 0;opacity:.9;}}
.k1::before{{background:linear-gradient(90deg,#6366f1,#818cf8)}}
.k2::before{{background:linear-gradient(90deg,#f59e0b,#fbbf24)}}
.k3::before{{background:linear-gradient(90deg,#3b82f6,#60a5fa)}}
.k4::before{{background:linear-gradient(90deg,#14b8a6,#2dd4bf)}}
.k5::before{{background:linear-gradient(90deg,#f97316,#fb923c)}}
.kpi::after{{content:'';position:absolute;bottom:-30px;right:-20px;width:80px;height:80px;border-radius:50%;opacity:.05;}}
.k1::after{{background:#6366f1}}.k2::after{{background:#f59e0b}}.k3::after{{background:#3b82f6}}.k4::after{{background:#14b8a6}}.k5::after{{background:#f97316}}
.kpi-v{{font-size:1.6rem;font-weight:800;letter-spacing:-.04em;line-height:1;font-family:'DM Mono',monospace;}}
.k1 .kpi-v{{color:#818cf8}}.k2 .kpi-v{{color:#fbbf24}}.k3 .kpi-v{{color:#60a5fa}}.k4 .kpi-v{{color:#2dd4bf}}.k5 .kpi-v{{color:#fb923c}}
.kpi-l{{font-size:.58rem;text-transform:uppercase;letter-spacing:1px;color:var(--muted);margin-top:6px;font-weight:700;font-family:'DM Mono',monospace;}}
.panel{{
  background:var(--bg2);border:1px solid var(--border);border-radius:14px;
  padding:22px 22px 18px;box-shadow:var(--card-shadow),var(--glow);
  transition:background .4s,border-color .35s;position:relative;overflow:hidden;
}}
.panel::after{{
  content:'';position:absolute;top:0;right:0;width:200px;height:200px;
  background:radial-gradient(circle at top right,rgba(20,184,166,0.04) 0%,transparent 60%);pointer-events:none;
}}
.panel-hd{{display:flex;align-items:flex-start;justify-content:space-between;margin-bottom:14px;flex-wrap:wrap;gap:10px;}}
.panel-ttl{{font-size:.65rem;font-weight:700;color:var(--text2);text-transform:uppercase;letter-spacing:1px;display:flex;align-items:center;gap:8px;font-family:'DM Mono',monospace;}}
.panel-dot{{width:7px;height:7px;border-radius:50%;flex-shrink:0;background:linear-gradient(135deg,#14b8a6,#6366f1);box-shadow:0 0 8px rgba(20,184,166,.5);}}
.leg{{display:flex;flex-wrap:wrap;gap:6px}}
.li{{
  display:flex;align-items:center;gap:5px;font-size:.66rem;color:var(--text2);font-weight:600;
  background:var(--bg3);border:1px solid var(--border);border-radius:6px;padding:4px 10px 4px 7px;
  transition:border-color .15s,background .4s;cursor:default;font-family:'DM Mono',monospace;letter-spacing:.3px;
}}
.li:hover{{border-color:rgba(20,184,166,.3);background:var(--bg4)}}
.ld{{width:12px;height:12px;border-radius:3px;flex-shrink:0;}}
.cw{{position:relative;width:100%;overflow-x:auto;margin-top:10px}}
canvas{{display:block;cursor:crosshair}}
#tt{{position:fixed;pointer-events:none;z-index:9999;opacity:0;transition:opacity .12s;min-width:260px;max-width:340px;}}
#tt.on{{opacity:1}}
#tt-body{{
  border-radius:10px;padding:13px 15px;font-family:'DM Mono',monospace;font-size:10.5px;
  border:1px solid rgba(20,184,166,0.25);box-shadow:0 10px 40px rgba(0,0,0,0.55);
}}
[data-theme="dark"] #tt-body{{background:#060c18;color:#dde4f0;}}
[data-theme="light"] #tt-body{{background:#fff;color:#1a2540;border-color:rgba(0,0,0,0.09);box-shadow:0 8px 30px rgba(0,0,0,0.13);}}
.tt-head{{font-size:11.5px;font-weight:700;margin-bottom:9px;padding-bottom:8px;border-bottom:1px solid rgba(255,255,255,0.08);}}
[data-theme="light"] .tt-head{{border-bottom-color:rgba(0,0,0,0.08);}}
.tt-ref{{color:#14b8a6;}}.tt-cmd{{color:var(--muted);font-size:9.5px;font-weight:400;margin-top:2px;}}
.tt-row{{display:flex;align-items:center;gap:8px;margin-bottom:7px;line-height:1.3;}}
.tt-row:last-child{{margin-bottom:0}}
.tt-swatch{{width:9px;height:9px;border-radius:2px;flex-shrink:0;}}
.tt-lbl{{flex:1;color:var(--text2);}}.tt-num{{font-weight:700;}}
.tt-pct{{color:var(--muted);font-size:9px;margin-left:4px;}}
.tt-encours{{margin-top:8px;padding:6px 10px;border-radius:7px;background:rgba(249,115,22,0.10);border:1px solid rgba(249,115,22,0.22);display:flex;align-items:center;gap:8px;font-size:10px;}}
.tt-encours-dot{{width:8px;height:8px;border-radius:50%;background:#f97316;flex-shrink:0;}}
.tt-encours-lbl{{flex:1;color:#fb923c;font-weight:600;}}.tt-encours-val{{font-weight:800;color:#f97316;font-size:11px;}}
.tt-sep{{border:none;border-top:1px solid rgba(255,255,255,0.06);margin:6px 0;}}
[data-theme="light"] .tt-sep{{border-top-color:rgba(0,0,0,0.07);}}
.pg-row{{margin-top:16px;padding-top:14px;border-top:1px solid var(--border);display:grid;grid-template-columns:repeat(auto-fit,minmax(180px,1fr));gap:12px;}}
.pi-lbl{{display:flex;justify-content:space-between;font-size:.65rem;margin-bottom:5px;font-family:'DM Mono',monospace;}}
.pi-lbl span:first-child{{color:var(--text2);font-weight:600;}}.pi-lbl .pct{{color:var(--muted);}}
.pt{{height:4px;background:var(--bg3);border-radius:99px;overflow:hidden;border:1px solid var(--border);}}
.pf{{height:100%;border-radius:99px;transition:width 1.3s cubic-bezier(.22,.68,0,1.2);width:0;}}
footer{{text-align:center;padding:16px;border-top:1px solid var(--border);}}
.footer-title{{font-size:.65rem;color:var(--muted);letter-spacing:.5px;font-family:'DM Mono',monospace;}}
.footer-title span{{color:var(--text2);}}
.footer-credit{{font-size:.48rem;color:var(--muted);letter-spacing:.8px;font-family:'DM Mono',monospace;margin-top:4px;opacity:.6;font-style:italic;}}
</style>
</head>
<body>
<div id="tt"><div id="tt-body"></div></div>
<header>
  <div class="logo-mark">RG</div>
  <h1>{titre_page}</h1>
  <span class="badge bd1">RÉCAP GLOBAL</span>
  <span class="badge bd2">Interactif</span>
  <button class="theme-btn" onclick="basculerTheme()">
    <div class="tog-track"><div class="tog-thumb"></div></div>
    <span id="theme-lbl">Mode clair</span>
  </button>
</header>
<div class="toolbar">
  <span class="tb-lbl">Vue</span>
  <div class="sw"><select id="sel" onchange="afficher(this.value)"></select></div>
  <span class="tb-ct"><b id="nb">0</b> graphe(s)</span>
</div>
<main id="main"></main>
<footer>
  <div class="footer-title"><span>{titre_page}</span></div>
  <div class="footer-credit">by Tsiory Burleigh</div>
</footer>
<script>
const TOUT={data_js};
const SEG_DEFS=[
  {{key:'seg_exp',rawKey:'total_exp',label:'Qtés expédiées',dark:'#14b8a6',light:'#0d9488',showZero:false}},
  {{key:'seg_livrees',rawKey:'qtes_livrees',label:'Qtés livrées',dark:'#3b82f6',light:'#2563eb',showZero:false}},
  {{key:'seg_dispatch',rawKey:'qtes_dispatchees',label:'Qtés dispatchées',dark:'#f59e0b',light:'#d97706',showZero:false}},
  {{key:'seg_reste',rawKey:null,label:'Non dispatché',dark:'#ef4444',light:'#dc2626',showZero:true}},
];
const KPI_SERIES=[
  {{key:'qte_commande',label:'QTÉ COMMANDÉE',kpi:'k1',color:'#818cf8'}},
  {{key:'qtes_dispatchees',label:'QTÉS DISPATCHÉES',kpi:'k2',color:'#fbbf24'}},
  {{key:'qtes_livrees',label:'QTÉS LIVRÉES',kpi:'k3',color:'#60a5fa'}},
  {{key:'total_exp',label:'QTÉS EXPÉDIÉES',kpi:'k4',color:'#2dd4bf'}},
  {{key:'qte_en_cours',label:'QTÉ EN COURS',kpi:'k5',color:'#fb923c'}},
];
const PG_SERIES_KEYS=['qtes_dispatchees','qtes_livrees','total_exp','qte_en_cours'];
let estSombre=true;
const GRAPHES={{}};const METAS={{}};
const ttEl=document.getElementById('tt');const ttBody=document.getElementById('tt-body');
const sc=seg=>estSombre?seg.dark:seg.light;
function attacherInfobulle(canvas,idx){{
  canvas.addEventListener('mousemove',e=>{{
    const g=GRAPHES[idx];const m=METAS[idx];if(!g||!m)return;
    const rect=canvas.getBoundingClientRect();
    const mx=e.clientX-rect.left;const my=e.clientY-rect.top;
    const aY=g.scales.y;const aX=g.scales.x;if(!aY)return;
    if(mx<aY.right||mx>aX.right||my<aY.top||my>aY.bottom){{cacherTT();return;}}
    const di=Math.round(aY.getValueForPixel(my));const d=m.d;
    if(di<0||di>=d.references.length){{cacherTT();return;}}
    const cb=aY.getPixelForValue(di);const dh=Math.abs(aY.getPixelForValue(0)-aY.getPixelForValue(1))*0.5*0.52;
    if(Math.abs(my-cb)>dh){{cacherTT();return;}}
    afficherTT(di,d,m.segs,e.clientX,e.clientY);
  }});
  canvas.addEventListener('mouseleave',cacherTT);
}}
function afficherTT(i,d,segs,cx,cy){{
  const ref=d.references[i];const cmd=d.qte_commande[i]||0;
  const disp=(d.qtes_dispatchees&&d.qtes_dispatchees[i])||0;
  const livr=(d.qtes_livrees&&d.qtes_livrees[i])||0;
  const enCours=(d.qte_en_cours&&d.qte_en_cours[i])||Math.max(0,disp-livr);
  let lignes='';
  segs.forEach(seg=>{{
    let qte=seg.rawKey===null?Math.max(0,cmd-disp):((d[seg.rawKey]&&d[seg.rawKey][i])||0);
    if(qte===0&&!seg.showZero)return;
    const pct=cmd>0?Math.round(qte/cmd*100):0;
    const sep=seg.rawKey===null?'<hr class="tt-sep"/>':'';
    lignes+=`${{sep}}<div class="tt-row"><div class="tt-swatch" style="background:${{sc(seg)}}"></div><span class="tt-lbl">${{seg.label}}</span><span class="tt-num">${{qte.toLocaleString('fr-FR')}}</span><span class="tt-pct">${{pct}}%</span></div>`;
  }});
  const ecp=cmd>0?Math.round(enCours/cmd*100):0;
  const bec=`<div class="tt-encours"><div class="tt-encours-dot"></div><span class="tt-encours-lbl">QTÉ EN COURS</span><span class="tt-encours-val">${{enCours.toLocaleString('fr-FR')}}</span><span class="tt-pct">${{ecp}}%</span></div>`;
  ttBody.innerHTML=`<div class="tt-head"><div class="tt-ref">${{ref}}</div><div class="tt-cmd">Commande: ${{cmd.toLocaleString('fr-FR')}}</div></div>${{lignes||'<div style="opacity:.5;font-size:9px">Aucune donnée</div>'}}${{bec}}`;
  const larg=ttEl.offsetWidth||280;const haut=ttEl.offsetHeight||150;
  const lv=window.innerWidth;const hv=window.innerHeight;
  let g=cx+16;let t=cy-haut/2;
  if(g+larg>lv-8)g=cx-larg-16;if(t<8)t=8;if(t+haut>hv-8)t=hv-haut-8;
  ttEl.style.left=g+'px';ttEl.style.top=t+'px';ttEl.classList.add('on');
}}
function cacherTT(){{ttEl.classList.remove('on');}}
function basculerTheme(){{
  estSombre=!estSombre;
  document.documentElement.setAttribute('data-theme',estSombre?'dark':'light');
  document.getElementById('theme-lbl').textContent=estSombre?'Mode clair':'Mode sombre';
  Object.keys(GRAPHES).forEach(idx=>{{GRAPHES[idx].destroy();delete GRAPHES[idx];construireGraphe(parseInt(idx));}});
}}
function init(){{
  const sel=document.getElementById('sel');const main=document.getElementById('main');
  if(!TOUT||!TOUT.length){{main.innerHTML='<p style="text-align:center;padding:80px;color:var(--muted)">Aucune donnée.</p>';return;}}
  document.getElementById('nb').textContent=TOUT.length;
  TOUT.forEach((bloc,idx)=>{{
    const opt=document.createElement('option');opt.value=idx;opt.textContent=bloc.titre;sel.appendChild(opt);
    const d=bloc.donnees;const segs=d.has_total_exp?SEG_DEFS.slice(0,4):SEG_DEFS.slice(1,4);
    METAS[idx]={{d,segs}};
    const kpiS=KPI_SERIES.filter(s=>s.key!=='total_exp'||d.has_total_exp);
    const tot={{}};kpiS.forEach(s=>{{tot[s.key]=(d[s.key]||[]).reduce((a,b)=>a+b,0);}});
    const totCmd=tot.qte_commande||1;
    const kpiH=kpiS.map(s=>`<div class="kpi ${{s.kpi}}"><div class="kpi-v">${{tot[s.key].toLocaleString('fr-FR')}}</div><div class="kpi-l">${{s.label}}</div></div>`).join('');
    const legH=segs.map(s=>`<div class="li"><div class="ld" style="background:${{sc(s)}}"></div>${{s.label}}</div>`).join('');
    const pgS=kpiS.filter(s=>PG_SERIES_KEYS.includes(s.key));
    const pgH=pgS.map(s=>{{const pct=Math.round(tot[s.key]/totCmd*100);return`<div><div class="pi-lbl"><span>${{s.label}}</span><span class="pct">${{pct}}%</span></div><div class="pt"><div class="pf" id="pf-${{idx}}-${{s.key}}" style="background:${{s.color}}"></div></div></div>`;}}).join('');
    const carte=document.createElement('div');carte.className='gc'+(idx===0?' active':'');carte.id='card-'+idx;
    carte.innerHTML=`<div class="gc-title">${{bloc.titre}}</div><div class="kpi-row">${{kpiH}}</div><div class="panel"><div class="panel-hd"><div class="panel-ttl"><div class="panel-dot"></div>Avancement par référence</div><div class="leg">${{legH}}</div></div><div class="cw"><canvas id="cv-${{idx}}"></canvas></div><div class="pg-row">${{pgH}}</div></div>`;
    main.appendChild(carte);
  }});
  construireGraphe(0);setTimeout(()=>animerBarres(0),450);
}}
function afficher(idx){{
  idx=parseInt(idx);cacherTT();
  document.querySelectorAll('.gc').forEach(c=>c.classList.remove('active'));
  const carte=document.getElementById('card-'+idx);
  if(carte){{carte.classList.add('active');construireGraphe(idx);setTimeout(()=>animerBarres(idx),140);}}
}}
function animerBarres(idx){{
  const d=TOUT[idx].donnees;const kpiS=KPI_SERIES.filter(s=>(s.key!=='total_exp'||d.has_total_exp)&&PG_SERIES_KEYS.includes(s.key));
  const cmd=(d.qte_commande||[]).reduce((a,b)=>a+b,0)||1;
  kpiS.forEach(s=>{{const t=(d[s.key]||[]).reduce((a,b)=>a+b,0);const el=document.getElementById(`pf-${{idx}}-${{s.key}}`);if(el)el.style.width=Math.min(100,Math.round(t/cmd*100))+'%';}});
}}
function construireGraphe(idx){{
  if(GRAPHES[idx])return;
  const d=TOUT[idx].donnees;const segs=METAS[idx].segs;if(!d)return;
  const canvas=document.getElementById('cv-'+idx);if(!canvas)return;
  const nb=d.references.length;canvas.style.height=Math.max(260,nb*36+80)+'px';
  const gridC=estSombre?'rgba(255,255,255,0.04)':'rgba(0,0,0,0.05)';
  const tickC=estSombre?'#3d5068':'#8fa3be';const axisC=estSombre?'#4a6080':'#6a80a0';
  const yTick=estSombre?'#7a8fab':'#4a6080';
  const xMax=Math.ceil(Math.max(...(d.qte_commande||[1]))*1.05);
  const jeux=segs.map((seg,si)=>({{label:seg.label,data:d[seg.key]||new Array(nb).fill(0),backgroundColor:sc(seg),borderWidth:0,borderRadius:si===segs.length-1?{{topRight:4,bottomRight:4,topLeft:0,bottomLeft:0}}:0,borderSkipped:true,barPercentage:0.52,categoryPercentage:0.78,stack:'principal'}}));
  GRAPHES[idx]=new Chart(canvas.getContext('2d'),{{type:'bar',data:{{labels:d.references,datasets:jeux}},options:{{indexAxis:'y',responsive:true,maintainAspectRatio:false,plugins:{{legend:{{display:false}},tooltip:{{enabled:false}}}},scales:{{x:{{stacked:true,max:xMax,grid:{{color:gridC,drawBorder:false,lineWidth:1}},border:{{display:false}},ticks:{{color:tickC,font:{{size:10,family:'DM Mono'}},callback:v=>v.toLocaleString('fr-FR'),maxTicksLimit:8}},title:{{display:true,text:'Quantités',color:axisC,font:{{size:10,weight:'600',family:'DM Mono'}},padding:{{top:8}}}}}},y:{{stacked:true,grid:{{display:false}},border:{{display:false}},ticks:{{color:yTick,font:{{size:10,weight:'600',family:'DM Mono'}},autoSkip:false,padding:8}},title:{{display:true,text:'RÉFÉRENCE',color:axisC,font:{{size:10,weight:'700',family:'DM Mono'}}}}}}}},animation:{{duration:0}}}}}});
  attacherInfobulle(canvas,idx);
}}
document.addEventListener('DOMContentLoaded',init);
</script>
</body>
</html>"""


# ══════════════════════════════════════════════════════════════════════════════
#                           MODULE ANALYSEUR PAR CLIENT
# ══════════════════════════════════════════════════════════════════════════════

class ModuleAnalyseurClient:

    @staticmethod
    def traiter_fichier_multi_clients(input_file, output_file, clients_selectionnes, log_callback):
        try:
            log_callback("=" * 70)
            log_callback("🚀 Début du traitement...")
            log_callback("📖 Lecture du fichier d'entrée...")
            df = pd.read_excel(input_file, sheet_name="SUIVI")

            colonnes_requises = [
                'CLIENT', 'REFERENCE', 'DATE DISPATCH', 'DATE LIVRAISON',
                'TYPE', 'TABLE', 'MATRICULE', 'PRENOM/ST', 'ETAT'
            ]
            for col in colonnes_requises:
                if col not in df.columns:
                    raise ValueError(f"Colonne '{col}' manquante.")

            df['DATE DISPATCH'] = pd.to_datetime(df['DATE DISPATCH'], errors='coerce')
            df['DATE LIVRAISON'] = pd.to_datetime(df['DATE LIVRAISON'], errors='coerce')

            log_callback("🔍 Filtrage des données (ETAT = PO ou SP)...")
            df = df[df['ETAT'].isin(['PO', 'SP'])].copy()
            log_callback(f"   ✅ {len(df)} lignes conservées")

            clients = sorted(df['CLIENT'].dropna().unique())
            if clients_selectionnes:
                clients = [c for c in clients if c in clients_selectionnes]

            if not clients:
                raise ValueError("Aucun client sélectionné ou trouvé.")

            log_callback(f"👥 Clients à traiter: {len(clients)}")
            donnees_feuilles = {}

            for client in clients:
                log_callback(f"\n📄 Traitement: {client}")
                df_client = df[df['CLIENT'] == client].copy()
                client_data = ModuleAnalyseurClient._generer_donnees_client(df_client, client)
                for sheet_name, data in client_data.items():
                    if sheet_name not in donnees_feuilles:
                        donnees_feuilles[sheet_name] = []
                    donnees_feuilles[sheet_name].append({'nom': client, 'data': data})

            log_callback("\n💾 Écriture du fichier Excel...")
            ModuleAnalyseurClient._ecrire_fichier_assemble(output_file, donnees_feuilles, log_callback)
            log_callback("=" * 70)
            log_callback(f"✅ Excel terminé! {len(clients)} client(s)")
            return len(clients)

        except Exception as e:
            raise Exception(f"Erreur de traitement: {str(e)}")

    @staticmethod
    def lire_clients(input_file):
        df = pd.read_excel(input_file, usecols=['CLIENT', 'ETAT'])
        df = df[df['ETAT'].isin(['PO', 'SP'])]
        return sorted(df['CLIENT'].dropna().unique().tolist())

    @staticmethod
    def _generer_donnees_client(df, client_name):
        recap_global = ModuleAnalyseurClient._creer_recap_global(df)
        recap_a = ModuleAnalyseurClient._creer_recap_type(df, 'A', 'INTERNE')
        recap_st = ModuleAnalyseurClient._creer_recap_type(df, 'ST', 'SOUS TRAITANT')
        recap_jour = ModuleAnalyseurClient._construire_tableau_periode(df, 'JOUR', "%d/%m/%Y")
        recap_semaine = ModuleAnalyseurClient._construire_tableau_periode(df, 'SEMAINE', "%d/%m/%Y")
        recap_mois = ModuleAnalyseurClient._construire_tableau_periode(df, 'MOIS', "%B %Y")
        recap_ref = ModuleAnalyseurClient._creer_recap_reference(df)
        recap_dispatch = ModuleAnalyseurClient._creer_recap_dispatch(df)
        recap_p = ModuleAnalyseurClient._creer_recap_personne(df)
        _, recap_temps_stats = ModuleAnalyseurClient._creer_recap_temps(df)
        recap_timing = ModuleAnalyseurClient._creer_timing_par_modele(df)

        a_et_st_combined = []
        if recap_a is not None and not recap_a.empty:
            a_et_st_combined.append(recap_a)
        if recap_st is not None and not recap_st.empty:
            a_et_st_combined.append(recap_st)
        a_et_st = pd.concat(a_et_st_combined, ignore_index=True) if a_et_st_combined else pd.DataFrame()

        return {
            "RECAP GLOBAL": recap_global,
            "A et ST": a_et_st,
            "RAMASSAGE PAR JOUR": recap_jour,
            "RAMASSAGE PAR SEMAINE": recap_semaine,
            "RAMASSAGE PAR MOIS": recap_mois,
            "RAMASSAGE PAR MODELES": recap_ref,
            "DISPATCH": recap_dispatch,
            "RAMASSAGE PAR PRESTATAIRE": recap_p,
            "STATS_TEMPS": recap_temps_stats,
            "TIMING PAR MODELE": recap_timing,
        }

    @staticmethod
    def _creer_recap_cumule_global(items):
        colonnes_somme = [
            'QTE COMMANDE', 'QTES DISPATCHEES', 'QTES LIVREES',
            'RESTE A DISPATCHER', 'QTES EN COURS'
        ]
        lignes = []
        colonnes_trouvees = set()

        for item in items:
            df = item['data']
            cols_map = {}
            for c in df.columns:
                c_upper = str(c).strip().upper()
                for cs in colonnes_somme:
                    if c_upper == cs:
                        cols_map[cs] = c
                        colonnes_trouvees.add(cs)

            total_row_data = None
            for idx, row in df.iterrows():
                first_val = str(row.iloc[0]).strip().upper() if pd.notna(row.iloc[0]) else ''
                if first_val == 'TOTAL':
                    total_row_data = row
                    break

            ligne = {'CLIENT': item['nom']}
            for cs in colonnes_somme:
                if cs in cols_map:
                    if total_row_data is not None:
                        val = total_row_data.get(cols_map[cs], 0)
                    else:
                        val = pd.to_numeric(df[cols_map[cs]], errors='coerce').sum()
                    try:
                        ligne[cs] = int(float(val)) if pd.notna(val) else 0
                    except (ValueError, TypeError):
                        ligne[cs] = 0
            lignes.append(ligne)

        if not lignes:
            return None

        colonnes_finales = ['CLIENT'] + [cs for cs in colonnes_somme if cs in colonnes_trouvees]
        df_recap = pd.DataFrame(lignes, columns=colonnes_finales)
        for col in colonnes_finales[1:]:
            df_recap[col] = df_recap[col].fillna(0).astype(int)

        total_general = {'CLIENT': 'TOTAL GENERAL'}
        for col in colonnes_finales[1:]:
            total_general[col] = int(df_recap[col].sum())
        df_recap.loc[len(df_recap)] = total_general
        return df_recap

    @staticmethod
    def _ecrire_fichier_assemble(output_file, donnees_feuilles, log_callback):
        recap_global_cumule = None
        if "RECAP GLOBAL" in donnees_feuilles:
            log_callback("📊 Calcul du TOTAL GENERAL...")
            recap_global_cumule = ModuleAnalyseurClient._creer_recap_cumule_global(
                donnees_feuilles["RECAP GLOBAL"]
            )

        with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
            for feuille, items in donnees_feuilles.items():
                row = 0
                if feuille == "RECAP GLOBAL" and recap_global_cumule is not None:
                    nb_cols = len(recap_global_cumule.columns)
                    titre_recap = pd.DataFrame(
                        [["TOTAL GENERAL"] + [''] * (nb_cols - 1)],
                        columns=recap_global_cumule.columns
                    )
                    titre_recap.to_excel(writer, sheet_name=feuille, startrow=row, index=False, header=False)
                    recap_global_cumule.to_excel(writer, sheet_name=feuille, startrow=row + 1, index=False)
                    row += len(recap_global_cumule) + 4
                    for item in items:
                        df = item['data']
                        if df is None or df.empty or len(df.columns) == 0:
                            continue
                        try:
                            titre_df = pd.DataFrame(
                                [[item['nom']] + [''] * (len(df.columns) - 1)],
                                columns=df.columns
                            )
                            titre_df.to_excel(writer, sheet_name=feuille, startrow=row, index=False, header=False)
                            df.to_excel(writer, sheet_name=feuille, startrow=row + 1, index=False)
                            row += len(df) + 3
                        except Exception as e:
                            log_callback(f"  ⚠️  Erreur '{feuille}' / '{item['nom']}': {e}")
                else:
                    for item in items:
                        df = item['data']
                        if df is None or df.empty or len(df.columns) == 0:
                            continue
                        try:
                            titre_df = pd.DataFrame(
                                [[item['nom']] + [''] * (len(df.columns) - 1)],
                                columns=df.columns
                            )
                            titre_df.to_excel(writer, sheet_name=feuille, startrow=row, index=False, header=False)
                            df.to_excel(writer, sheet_name=feuille, startrow=row + 1, index=False)
                            row += len(df) + 3
                        except Exception as e:
                            log_callback(f"  ⚠️  Erreur '{feuille}' / '{item['nom']}': {e}")

        log_callback("🎨 Application du style...")
        StylesExcel.init_random_assembleur_colors()
        MiseEnForme.appliquer(output_file)

    @staticmethod
    def _creer_recap_global(df):
        grouped = df.groupby('REFERENCE')
        qte_commande = df[df['ETAT'] == 'PO'].groupby('REFERENCE').size()
        recap = pd.DataFrame({
            'REFERENCE': grouped.size().index,
            'QTE COMMANDE': qte_commande.reindex(grouped.size().index, fill_value=0).values,
            'QTES DISPATCHEES': grouped['DATE DISPATCH'].apply(lambda x: x.notna().sum()),
            'QTES LIVREES': grouped['DATE LIVRAISON'].apply(lambda x: x.notna().sum())
        })
        recap['RESTE A DISPATCHER'] = (recap['QTE COMMANDE'] - recap['QTES DISPATCHEES']).clip(lower=0)
        recap['QTES EN COURS'] = recap['QTES DISPATCHEES'] - recap['QTES LIVREES']
        total = pd.DataFrame({
            col: ['TOTAL'] if col == 'REFERENCE' else [recap[col].sum()]
            for col in recap.columns
        })
        return pd.concat([recap, total], ignore_index=True)

    @staticmethod
    def _creer_recap_type(df, type_filtre, nom_colonne):
        df_filtre = df[df['TYPE'] == type_filtre]
        if len(df_filtre) == 0:
            return pd.DataFrame()
        if type_filtre == 'ST':
            grouped = df_filtre.groupby(['MATRICULE', 'REFERENCE'])
            records = []
            for (matricule, reference), group in grouped:
                date_dispatch = (
                    group['DATE DISPATCH'].dropna().min()
                    if group['DATE DISPATCH'].notna().any() else pd.NaT
                )
                qtes_dispatchees = group['DATE DISPATCH'].notna().sum()
                qtes_livrees = group['DATE LIVRAISON'].notna().sum()
                records.append({
                    'SOUS TRAITANT': matricule, 'REFERENCE': reference,
                    'DATE DISPATCH': date_dispatch.strftime('%d/%m/%Y') if pd.notna(date_dispatch) else '',
                    'QTES DISPATCHEES': qtes_dispatchees, 'QTES LIVREES': qtes_livrees,
                    'QTES EN COURS': qtes_dispatchees - qtes_livrees
                })
            recap = pd.DataFrame(records)
            if len(recap) > 0:
                recap = recap.sort_values(['SOUS TRAITANT', 'REFERENCE']).reset_index(drop=True)
            total = pd.DataFrame({
                'SOUS TRAITANT': ['TOTAL'], 'REFERENCE': [''], 'DATE DISPATCH': [''],
                'QTES DISPATCHEES': [recap['QTES DISPATCHEES'].sum() if len(recap) > 0 else 0],
                'QTES LIVREES': [recap['QTES LIVREES'].sum() if len(recap) > 0 else 0],
                'QTES EN COURS': [recap['QTES EN COURS'].sum() if len(recap) > 0 else 0],
            })
            return pd.concat([recap, total], ignore_index=True)
        else:
            grouped = df_filtre.groupby('REFERENCE')
            recap = pd.DataFrame({
                nom_colonne: grouped.size().index,
                'QTES DISPATCHEES': grouped['DATE DISPATCH'].apply(lambda x: x.notna().sum()),
                'QTES LIVREES': grouped['DATE LIVRAISON'].apply(lambda x: x.notna().sum())
            })
            recap['QTES EN COURS'] = recap['QTES DISPATCHEES'] - recap['QTES LIVREES']
            recap = recap.sort_values(nom_colonne).reset_index(drop=True)
            total = pd.DataFrame({
                nom_colonne: ['TOTAL'],
                'QTES DISPATCHEES': [recap['QTES DISPATCHEES'].sum()],
                'QTES LIVREES': [recap['QTES LIVREES'].sum()],
                'QTES EN COURS': [recap['QTES EN COURS'].sum()]
            })
            return pd.concat([recap, total], ignore_index=True)

    @staticmethod
    def _construire_tableau_periode(df, periode, format_date):
        df_copy = df.copy()
        if periode == 'JOUR':
            df_copy['PERIODE'] = df_copy['DATE LIVRAISON'].dt.date
        elif periode == 'SEMAINE':
            df_copy['PERIODE'] = df_copy['DATE LIVRAISON'].dt.to_period('W').dt.start_time.dt.date
        else:
            df_copy['PERIODE'] = df_copy['DATE LIVRAISON'].dt.to_period('M').dt.start_time.dt.date

        periodes_uniques = sorted(df_copy['PERIODE'].dropna().unique())
        labels = [
            p.strftime(format_date).capitalize() if periode == 'MOIS' else p.strftime(format_date)
            for p in periodes_uniques
        ]
        colonnes = ['', 'DISPATCHEES'] + labels + ['TOTAL LIVREES', 'RESTE A LIVRER']
        recap = pd.DataFrame(columns=colonnes)

        def ajouter_ligne(label, data):
            dis = data['DATE DISPATCH'].notna().sum()
            liv = data['DATE LIVRAISON'].notna().sum()
            counts = [(data['PERIODE'] == p).sum() for p in periodes_uniques]
            recap.loc[len(recap)] = [label, dis] + counts + [liv, dis - liv]

        df_A = df_copy[df_copy['TYPE'] == 'A']
        if len(df_A) > 0:
            ajouter_ligne('A', df_A)
            for t in sorted(df_copy[df_copy['TYPE'] == 'A']['TABLE'].dropna().unique()):
                df_t = df_copy[(df_copy['TABLE'] == t) & (df_copy['TYPE'] == 'A')]
                if len(df_t) > 0:
                    ajouter_ligne(t, df_t)

        df_ST = df_copy[df_copy['TYPE'] == 'ST']
        if len(df_ST) > 0:
            ajouter_ligne('ST', df_ST)
            for m in sorted(df_ST['MATRICULE'].dropna().unique()):
                df_m = df_ST[df_ST['MATRICULE'] == m]
                if len(df_m) > 0:
                    ajouter_ligne(str(m), df_m)

        ajouter_ligne('TOTAL', df_copy)
        return recap

    @staticmethod
    def _creer_recap_reference(df):
        df_ref = df[df['DATE DISPATCH'].notna()].copy()
        dates = sorted(df_ref['DATE LIVRAISON'].dropna().dt.date.unique())
        colonnes = (
            ['REFERENCE', 'DISPATCHEES']
            + [d.strftime("%d/%m/%Y") for d in dates]
            + ['TOTAL LIVREES', 'RESTE A LIVRER']
        )
        recap = pd.DataFrame(columns=colonnes)
        for ref, group in df_ref.groupby('REFERENCE'):
            livraisons = group['DATE LIVRAISON'].dt.date.value_counts().to_dict()
            ligne = {
                'REFERENCE': ref, 'DISPATCHEES': len(group),
                'TOTAL LIVREES': group['DATE LIVRAISON'].notna().sum(),
                'RESTE A LIVRER': len(group) - group['DATE LIVRAISON'].notna().sum()
            }
            for d in dates:
                ligne[d.strftime("%d/%m/%Y")] = livraisons.get(d, 0)
            recap.loc[len(recap)] = ligne
        total = {'REFERENCE': 'TOTAL'}
        for col in colonnes[1:]:
            total[col] = recap[col].sum() if col in recap.columns else 0
        recap.loc[len(recap)] = total
        return recap[colonnes]

    @staticmethod
    def _creer_recap_dispatch(df):
        dates = sorted(df['DATE DISPATCH'].dropna().dt.date.unique())
        colonnes = (
            ['REFERENCE', 'QTE COMMANDE']
            + [d.strftime("%d/%m/%Y") for d in dates]
            + ['TOTAL DISPATCHEES', 'RESTE A DISPATCHER']
        )
        recap = pd.DataFrame(columns=colonnes)
        for ref, group in df.groupby('REFERENCE'):
            dispatch = group['DATE DISPATCH'].dt.date.value_counts().to_dict()
            qte_commande = len(group[group['ETAT'] == 'PO'])
            ligne = {
                'REFERENCE': ref, 'QTE COMMANDE': qte_commande,
                'TOTAL DISPATCHEES': group['DATE DISPATCH'].notna().sum(),
                'RESTE A DISPATCHER': max(0, qte_commande - group['DATE DISPATCH'].notna().sum())
            }
            for d in dates:
                ligne[d.strftime("%d/%m/%Y")] = dispatch.get(d, 0)
            recap.loc[len(recap)] = ligne
        total = {'REFERENCE': 'TOTAL'}
        for col in colonnes[1:]:
            total[col] = recap[col].sum() if col in recap.columns else 0
        recap.loc[len(recap)] = total
        return recap[colonnes]

    @staticmethod
    def _creer_recap_personne(df):
        df_p = df[df['DATE DISPATCH'].notna()].copy()
        dates = sorted(df_p['DATE LIVRAISON'].dropna().dt.date.unique())
        colonnes = (
            ['MATRICULE', 'DETAILS RESTE A LIVRER', 'TABLE', 'DISPATCHEES']
            + [d.strftime("%d/%m/%Y") for d in dates]
            + ['TOTAL LIVREES', 'RESTE A LIVRER']
        )
        recap = pd.DataFrame(columns=colonnes)
        for (mat, table), group in df_p.groupby(['MATRICULE', 'TABLE']):
            livraisons = group['DATE LIVRAISON'].dt.date.value_counts().to_dict()
            refs_non_livrees = group[group['DATE LIVRAISON'].isna()]['REFERENCE'].tolist()
            ligne = {
                'MATRICULE': mat,
                'DETAILS RESTE A LIVRER': ', '.join([str(r) for r in refs_non_livrees if pd.notna(r)]),
                'TABLE': table, 'DISPATCHEES': len(group),
                'TOTAL LIVREES': group['DATE LIVRAISON'].notna().sum(),
                'RESTE A LIVRER': len(group) - group['DATE LIVRAISON'].notna().sum()
            }
            for d in dates:
                ligne[d.strftime("%d/%m/%Y")] = livraisons.get(d, 0)
            recap.loc[len(recap)] = ligne
        return recap[colonnes]

    @staticmethod
    def _creer_timing_par_modele(df):
        empty_cols = ['REFERENCE', 'NB PIECES', 'MOYENNE (h)', 'MEDIANE (h)', 'MOYENNE (j)', 'MEDIANE (j)']
        empty = pd.DataFrame(columns=empty_cols)
        col_timing = next((col for col in df.columns if col.strip().upper() == 'TIMING'), None)
        if col_timing is None:
            return empty
        df_t = df[df['DATE LIVRAISON'].notna()].copy()
        df_t['TIMING_VAL'] = pd.to_numeric(df_t[col_timing], errors='coerce')
        df_t = df_t[df_t['TIMING_VAL'].notna() & (df_t['TIMING_VAL'] > 0)]
        if len(df_t) == 0:
            return empty
        rows = []
        for ref, group in df_t.groupby('REFERENCE'):
            vals_h = group['TIMING_VAL'].dropna().tolist()
            vals_j = [v / 8 for v in vals_h]
            if not vals_h:
                continue
            rows.append({
                'REFERENCE': ref, 'NB PIECES': len(vals_h),
                'MOYENNE (h)': round(sum(vals_h) / len(vals_h), 2),
                'MEDIANE (h)': round(statistics.median(vals_h), 2),
                'MOYENNE (j)': round(sum(vals_j) / len(vals_j), 2),
                'MEDIANE (j)': round(statistics.median(vals_j), 2),
            })
        if not rows:
            return empty
        recap = pd.DataFrame(rows, columns=empty_cols)
        return recap.sort_values('REFERENCE').reset_index(drop=True)

    @staticmethod
    def _creer_recap_temps(df):
        empty_stats = pd.DataFrame(columns=[
            'REFERENCE', 'NB PIECES',
            'MOY HEURES', 'MEDIANE HEURES', 'MOY TRONQUEE HEURES',
            'MOY JOURS', 'MEDIANE JOURS', 'MOY TRONQUEE JOURS'
        ])
        col_heures = None
        col_jours = None
        for col in df.columns:
            col_up = col.upper()
            if 'TEMPS' in col_up and 'ECOULE' in col_up:
                col_heures = col
            if 'DATE' in col_up and 'ECOULE' in col_up:
                col_jours = col
        if col_heures is None or col_jours is None:
            return None, empty_stats
        df_temps = df[df['DATE LIVRAISON'].notna()].copy()
        df_temps['HEURES'] = pd.to_numeric(df_temps[col_heures], errors='coerce')
        df_temps['JOURS'] = pd.to_numeric(df_temps[col_jours], errors='coerce')
        df_temps = df_temps[df_temps['HEURES'].notna() & df_temps['JOURS'].notna()]
        if len(df_temps) == 0:
            return None, empty_stats

        def moyenne_tronquee(serie, pct=0.10):
            n = len(serie)
            k = math.floor(n * pct)
            if k == 0 or n - 2 * k <= 0:
                return serie.mean()
            return serie.sort_values().values[k: n - k].mean()

        stats_rows = []
        for ref, group in df_temps.groupby('REFERENCE'):
            heures = group['HEURES'].dropna()
            jours = group['JOURS'].dropna()
            if len(heures) == 0:
                continue
            stats_rows.append({
                'REFERENCE': ref, 'NB PIECES': len(heures),
                'MOY HEURES': round(heures.mean(), 2),
                'MEDIANE HEURES': round(heures.median(), 2),
                'MOY TRONQUEE HEURES': round(moyenne_tronquee(heures), 2),
                'MOY JOURS': round(jours.mean(), 2),
                'MEDIANE JOURS': round(jours.median(), 2),
                'MOY TRONQUEE JOURS': round(moyenne_tronquee(jours), 2),
            })
        recap_stats = pd.DataFrame(stats_rows, columns=[
            'REFERENCE', 'NB PIECES',
            'MOY HEURES', 'MEDIANE HEURES', 'MOY TRONQUEE HEURES',
            'MOY JOURS', 'MEDIANE JOURS', 'MOY TRONQUEE JOURS'
        ])
        return None, recap_stats


# ══════════════════════════════════════════════════════════════════════════════
#                           APPLICATION FUSIONNÉE
# ══════════════════════════════════════════════════════════════════════════════

class ExcelProcessorApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Générateur Multi-Clients — Excel + Graphes HTML")
        self.root.geometry("980x820")
        self.root.minsize(820, 700)

        self.input_file = None
        self.output_xlsx = None
        self.output_html = None
        self.client_vars = {}

        self.create_widgets()

    def create_widgets(self):
        canvas = tk.Canvas(self.root, highlightthickness=0)
        scrollbar = ttk.Scrollbar(self.root, orient="vertical", command=canvas.yview)
        self.scroll_frame = ttk.Frame(canvas)
        self.scroll_frame.bind(
            "<Configure>",
            lambda e: canvas.configure(scrollregion=canvas.bbox("all"))
        )
        canvas.create_window((0, 0), window=self.scroll_frame, anchor="nw")
        canvas.configure(yscrollcommand=scrollbar.set)
        canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        main_frame = ttk.Frame(self.scroll_frame, padding=25)
        main_frame.pack(fill=BOTH, expand=YES)

        # ── Titre ──
        ttk.Label(
            main_frame,
            text="📊 Générateur de Rapport Multi-Clients",
            font=("Segoe UI", 18, "bold"),
            bootstyle="primary"
        ).pack(pady=(10, 3))

        ttk.Label(
            main_frame,
            text="Génère simultanément le fichier Excel (.xlsx) et les graphes interactifs (.html)",
            font=("Segoe UI", 10),
            bootstyle="secondary"
        ).pack(pady=(0, 15))

        # ── Fichier source ──
        input_frame = ttk.Labelframe(
            main_frame, text="📁 Fichier source (SUIVI)", padding=18, bootstyle="info"
        )
        input_frame.pack(fill=X, pady=8)
        self.input_label = ttk.Label(
            input_frame, text="Aucun fichier sélectionné",
            font=("Segoe UI", 9), bootstyle="secondary"
        )
        self.input_label.pack(side=LEFT, fill=X, expand=YES, padx=5)
        ttk.Button(
            input_frame, text="📂 Parcourir",
            command=self.select_input_file,
            bootstyle="info", width=15
        ).pack(side=RIGHT, padx=5)

        # ── Sélection clients ──
        self.clients_frame = ttk.Labelframe(
            main_frame, text="👥 Clients à générer", padding=14, bootstyle="warning"
        )
        self.clients_frame.pack(fill=X, pady=8)
        self.clients_inner = ttk.Frame(self.clients_frame)
        self.clients_inner.pack(fill=X)
        self.clients_placeholder = ttk.Label(
            self.clients_inner,
            text="← Chargez un fichier source pour voir les clients",
            font=("Segoe UI", 9, "italic"), bootstyle="secondary"
        )
        self.clients_placeholder.pack(pady=5)
        self.btn_frame = ttk.Frame(self.clients_frame)
        self.btn_frame.pack(fill=X, pady=(5, 0))
        self.btn_all = ttk.Button(
            self.btn_frame, text="✅ Tout sélectionner",
            command=self.select_all_clients,
            bootstyle="success-outline", width=20
        )
        self.btn_none = ttk.Button(
            self.btn_frame, text="☐ Tout désélectionner",
            command=self.deselect_all_clients,
            bootstyle="secondary-outline", width=22
        )
        self.btn_all.pack_forget()
        self.btn_none.pack_forget()

        # ── Fichiers de sortie (côte à côte) ──
        sortie_frame = ttk.Frame(main_frame)
        sortie_frame.pack(fill=X, pady=8)
        sortie_frame.columnconfigure(0, weight=1)
        sortie_frame.columnconfigure(1, weight=1)

        # Excel
        xlsx_frame = ttk.Labelframe(
            sortie_frame, text="💾 Fichier Excel (.xlsx)", padding=14, bootstyle="success"
        )
        xlsx_frame.grid(row=0, column=0, sticky=NSEW, padx=(0, 6))
        self.xlsx_label = ttk.Label(
            xlsx_frame, text="Non défini",
            font=("Segoe UI", 9), bootstyle="secondary"
        )
        self.xlsx_label.pack(side=LEFT, fill=X, expand=YES, padx=5)
        ttk.Button(
            xlsx_frame, text="💾 Parcourir",
            command=self.select_output_xlsx,
            bootstyle="success", width=13
        ).pack(side=RIGHT)

        # HTML
        html_frame = ttk.Labelframe(
            sortie_frame, text="🌐 Graphes HTML (.html)", padding=14, bootstyle="primary"
        )
        html_frame.grid(row=0, column=1, sticky=NSEW, padx=(6, 0))
        self.html_label = ttk.Label(
            html_frame, text="Non défini",
            font=("Segoe UI", 9), bootstyle="secondary"
        )
        self.html_label.pack(side=LEFT, fill=X, expand=YES, padx=5)
        ttk.Button(
            html_frame, text="🌐 Parcourir",
            command=self.select_output_html,
            bootstyle="primary", width=13
        ).pack(side=RIGHT)

        # ── Titre du graphe ──
        titre_frame = ttk.Labelframe(
            main_frame, text="✏️ Titre du rapport HTML", padding=14, bootstyle="secondary"
        )
        titre_frame.pack(fill=X, pady=8)
        self.titre_var = tk.StringVar(value="Suivi de production — Graphes")
        ttk.Entry(
            titre_frame, textvariable=self.titre_var,
            font=("Segoe UI", 10)
        ).pack(fill=X, padx=5)

        # ── Option ouvrir navigateur ──
        self.ouvrir_nav_var = BooleanVar(value=True)
        ttk.Checkbutton(
            main_frame,
            text="🌐  Ouvrir le fichier HTML dans le navigateur après génération",
            variable=self.ouvrir_nav_var,
            bootstyle="primary-round-toggle"
        ).pack(pady=6)

        # ── Bouton principal ──
        ttk.Button(
            main_frame,
            text="⚡  Générer Excel + Graphes HTML",
            command=self.process_file,
            bootstyle="danger",
            width=40
        ).pack(pady=20, ipady=6)

        # ── Journal ──
        log_frame = ttk.Labelframe(
            main_frame, text="📋 Journal d'activité", padding=14, bootstyle="secondary"
        )
        log_frame.pack(fill=BOTH, expand=YES, pady=8)
        log_scroll = ttk.Scrollbar(log_frame, orient="vertical")
        log_scroll.pack(side=RIGHT, fill=Y)
        self.log_text = ttk.Text(
            log_frame, height=12, width=80,
            font=("Consolas", 9),
            yscrollcommand=log_scroll.set
        )
        self.log_text.pack(fill=BOTH, expand=YES)
        log_scroll.config(command=self.log_text.yview)

    # ── Helpers ─────────────────────────────────────────────────────────────

    def log(self, message):
        self.log_text.insert(END, message + "\n")
        self.log_text.see(END)
        self.root.update()

    def select_input_file(self):
        filename = filedialog.askopenfilename(
            title="Fichier Excel source",
            filetypes=[("Excel files", "*.xlsx *.xls")]
        )
        if filename:
            self.input_file = filename
            self.input_label.config(text=os.path.basename(filename))
            self.log(f"✓ Source : {filename}")
            # Auto-remplir les sorties
            base = os.path.splitext(filename)[0]
            if not self.output_xlsx:
                self.output_xlsx = base + "_rapport.xlsx"
                self.xlsx_label.config(text=os.path.basename(self.output_xlsx))
            if not self.output_html:
                self.output_html = base + "_graphes.html"
                self.html_label.config(text=os.path.basename(self.output_html))
            nom = os.path.splitext(os.path.basename(filename))[0]
            self.titre_var.set(f"Suivi — {nom}")
            self._charger_clients()

    def _charger_clients(self):
        try:
            clients = ModuleAnalyseurClient.lire_clients(self.input_file)
        except Exception as e:
            self.log(f"⚠️  Lecture clients: {e}")
            return
        for widget in self.clients_inner.winfo_children():
            widget.destroy()
        self.client_vars.clear()
        if not clients:
            ttk.Label(self.clients_inner, text="Aucun client trouvé.", bootstyle="danger").pack()
            return
        self.clients_placeholder.pack_forget()
        cols = 4
        for idx, client in enumerate(clients):
            var = BooleanVar(value=True)
            self.client_vars[client] = var
            ttk.Checkbutton(
                self.clients_inner, text=str(client), variable=var,
                bootstyle="primary-round-toggle"
            ).grid(row=idx // cols, column=idx % cols, sticky=W, padx=10, pady=3)
        self.btn_all.pack(side=LEFT, padx=5)
        self.btn_none.pack(side=LEFT, padx=5)
        self.log(f"👥 {len(clients)} client(s): {', '.join(str(c) for c in clients)}")

    def select_all_clients(self):
        for var in self.client_vars.values():
            var.set(True)

    def deselect_all_clients(self):
        for var in self.client_vars.values():
            var.set(False)

    def select_output_xlsx(self):
        filename = filedialog.asksaveasfilename(
            title="Fichier Excel de sortie",
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")]
        )
        if filename:
            self.output_xlsx = filename
            self.xlsx_label.config(text=os.path.basename(filename))
            self.log(f"✓ Excel : {filename}")

    def select_output_html(self):
        filename = filedialog.asksaveasfilename(
            title="Fichier HTML de sortie",
            defaultextension=".html",
            filetypes=[("HTML files", "*.html")]
        )
        if filename:
            self.output_html = filename
            self.html_label.config(text=os.path.basename(filename))
            self.log(f"✓ HTML : {filename}")

    def process_file(self):
        if not self.input_file:
            messagebox.showerror("Erreur", "Sélectionnez un fichier source.")
            return
        if not self.output_xlsx:
            messagebox.showerror("Erreur", "Définissez le fichier Excel de sortie.")
            return
        if not self.output_html:
            messagebox.showerror("Erreur", "Définissez le fichier HTML de sortie.")
            return

        clients_selectionnes = [c for c, v in self.client_vars.items() if v.get()]
        if not clients_selectionnes:
            messagebox.showwarning("Attention", "Sélectionnez au moins un client.")
            return

        erreurs = []

        # ── ÉTAPE 1 : Excel ──────────────────────────────────────────────
        try:
            nb_clients = ModuleAnalyseurClient.traiter_fichier_multi_clients(
                self.input_file, self.output_xlsx, clients_selectionnes, self.log
            )
        except Exception as e:
            self.log(f"❌ Erreur Excel : {e}")
            erreurs.append(f"Excel : {e}")
            nb_clients = 0

        # ── ÉTAPE 2 : HTML (depuis le xlsx généré) ───────────────────────
        html_ok = False
        if nb_clients > 0:
            try:
                self.log("\n🌐 Génération du fichier HTML...")
                blocs = LecteurRecapGlobal.lire_fichier(self.output_xlsx)
                self.log(f"   → {len(blocs)} bloc(s) détecté(s)")
                GenerateurHTML.generer(blocs, self.output_html, self.titre_var.get())
                self.log(f"   ✅ {os.path.basename(self.output_html)} créé !")
                html_ok = True
            except Exception as e:
                self.log(f"❌ Erreur HTML : {e}")
                erreurs.append(f"HTML : {e}")

        # ── Résumé ───────────────────────────────────────────────────────
        if erreurs:
            messagebox.showerror("Erreurs", "\n".join(erreurs))
        else:
            resume = (
                f"✅ Génération terminée !\n\n"
                f"👥 {nb_clients} client(s) traité(s)\n\n"
                f"📊 Excel  →  {os.path.basename(self.output_xlsx)}\n"
                f"🌐 HTML   →  {os.path.basename(self.output_html)}"
            )
            messagebox.showinfo("Succès", resume)

        if html_ok and self.ouvrir_nav_var.get():
            webbrowser.open("file:///" + os.path.abspath(self.output_html))


# ══════════════════════════════════════════════════════════════════════════════

def main():
    root = ttk.Window(themename="cosmo")
    app = ExcelProcessorApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()